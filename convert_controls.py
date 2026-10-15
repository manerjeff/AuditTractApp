import json
import os
import openpyxl

def convert_excel_to_json(excel_path, output_json_path):
    if not os.path.exists(excel_path):
        print(f"⚠️ Excel file not found at {excel_path}. Skipping compilation.")
        return

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    controls = []

    # Process NIST Sheet
    if "NIST" in wb.sheetnames or "Sheet1" in wb.sheetnames:
        sheet = wb["NIST"] if "NIST" in wb.sheetnames else wb.active
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            controls.append({
                "id": str(row[0]).strip(),
                "frameworkKey": "NIST",
                "title": str(row[1]).strip() if len(row) > 1 and row[1] else "Untitled Control",
                "category": str(row[2]).strip() if len(row) > 2 and row[2] else "General",
                "controlStatement": str(row[3]).strip() if len(row) > 3 and row[3] else None,
                "discussion": str(row[4]).strip() if len(row) > 4 and row[4] else None,
                "relatedControls": [r.strip() for r in str(row[5]).split(",")] if len(row) > 5 and row[5] else []
            })

    # Process SOC 2 Sheet
    if "SOC2" in wb.sheetnames:
        sheet = wb["SOC2"]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            controls.append({
                "id": str(row[0]).strip(),
                "frameworkKey": "SOC2",
                "title": str(row[1]).strip() if len(row) > 1 and row[1] else "Untitled TSC Control",
                "category": str(row[2]).strip() if len(row) > 2 and row[2] else "Trust Services Criteria",
                "controlStatement": str(row[3]).strip() if len(row) > 3 and row[3] else None,
                "discussion": str(row[4]).strip() if len(row) > 4 and row[4] else None,
                "relatedControls": []
            })

    output = {
        "version": "1.0.0",
        "controls": controls
    }

    with open(output_json_path, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=2)

    print(f"🎉 Successfully compiled {len(controls)} controls into {output_json_path}")

if __name__ == "__main__":
    # Change "controls.json" to write directly inside the Services & Pipeline directory
    convert_excel_to_json("NIST_SP_800_53.xlsx", "AuditTractApp/Services & Pipeline/controls.json")

